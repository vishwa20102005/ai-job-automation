"""Document generation: PDF and DOCX from text content."""
import io


def generate_pdf(content: str, title: str = "Document") -> bytes:
    """Generate PDF from text content using ReportLab."""
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )

    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Heading1"],
        fontSize=16,
        spaceAfter=12,
    )
    story.append(Paragraph(title, title_style))
    story.append(Spacer(1, 12))

    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontSize=11,
        leading=16,
        alignment=TA_LEFT,
    )

    for para in content.split("\n\n"):
        para = para.strip()
        if para:
            safe = para.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            # Handle line breaks within paragraph
            safe = safe.replace("\n", "<br/>")
            story.append(Paragraph(safe, body_style))
            story.append(Spacer(1, 8))

    doc.build(story)
    return buffer.getvalue()


def generate_docx(content: str, title: str = "Document") -> bytes:
    """Generate DOCX from text content using python-docx."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    # Title
    heading = doc.add_heading(title, level=0)
    for run in heading.runs:
        run.font.size = Pt(16)

    # Body
    for para in content.split("\n\n"):
        para = para.strip()
        if para:
            doc.add_paragraph(para)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
